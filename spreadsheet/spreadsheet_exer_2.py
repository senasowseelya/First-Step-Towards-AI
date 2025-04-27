from openpyxl.reader.excel import load_workbook

inventory_workbook = load_workbook("inventory.xlsx")
products_list = inventory_workbook["Sheet1"]
#print(products_list)   #<Worksheet "Sheet1">

products_per_supplier={}

for product_row in range(2,products_list.max_row+1):
    #print(product_row)
    supplier_name = products_list.cell(product_row,4).value  # because supplier name is at column 4
    inventory = products_list.cell(product_row,2).value
    price = products_list.cell(product_row,3).value
    #print(supplier_name)
    if supplier_name not in products_per_supplier:
        products_per_supplier[supplier_name] = (price*inventory)
    else:
        products_per_supplier[supplier_name] += (price*inventory)

print(products_per_supplier)